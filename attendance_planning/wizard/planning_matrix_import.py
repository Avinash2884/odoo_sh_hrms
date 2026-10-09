# -*- coding: utf-8 -*-
from odoo import api, fields, models
from odoo.exceptions import UserError
import base64
import io
from datetime import datetime
import pytz
import logging

try:
    import openpyxl
except ImportError:
    openpyxl = None

_logger = logging.getLogger(__name__)


class PlanningMatrixImport(models.TransientModel):
    _name = 'planning.matrix.import'
    _description = 'Import Planning Matrix'

    upload_file = fields.Binary(string='Upload Excel Roster', required=True)
    file_name = fields.Char(string='File Name')

    # ── Holiday warning state ──
    holiday_warning_ids = fields.One2many(
        'planning.matrix.import.warning',
        'wizard_id',
        string='Holiday Conflicts'
    )
    confirmed = fields.Boolean(default=False)

    def _get_holiday_map(self):
        """Returns {date: holiday_name} for all public holidays."""
        tz_name = self.env.company.resource_calendar_id.tz or 'Asia/Kolkata'
        local_tz = pytz.timezone(tz_name)
        leaves = self.env['resource.calendar.leaves'].search([
            ('resource_id', '=', False),
            '|',
            ('company_id', '=', self.env.company.id),
            ('company_id', '=', False),
        ])
        holiday_map = {}
        for leaf in leaves:
            if leaf.date_from:
                utc_dt = leaf.date_from.replace(tzinfo=pytz.utc)
                local_date = utc_dt.astimezone(local_tz).date()
                holiday_map[local_date] = leaf.name or 'Public Holiday'
        return holiday_map

    def _parse_excel(self):
        """Parse Excel and return slots_to_create list."""
        if not openpyxl:
            raise UserError("The 'openpyxl' Python library is missing.")
        if not self.upload_file:
            raise UserError("Please upload a file.")

        file_data = base64.b64decode(self.upload_file)
        data = io.BytesIO(file_data)
        try:
            wb = openpyxl.load_workbook(data, data_only=True)
            sheet = wb.active
        except Exception as e:
            raise UserError(f"Could not read the Excel file. Error: {str(e)}")

        headers = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        date_map = {}
        for col_idx in range(3, len(headers)):
            val = headers[col_idx]
            if not val:
                continue
            if isinstance(val, datetime):
                date_map[col_idx] = val.date()
            else:
                try:
                    date_map[col_idx] = datetime.strptime(str(val).strip(), '%m/%d/%Y').date()
                except ValueError:
                    continue

        if not date_map:
            raise UserError("Could not find any valid dates in the header row.")

        slots_to_create = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            emp_code = row[0]
            if not emp_code:
                continue
            emp_code = str(emp_code).strip()

            for col_idx, target_date in date_map.items():
                shift_code = row[col_idx]

                if not shift_code or str(shift_code).strip().upper() in ['WO', 'OFF', '']:
                    slots_to_create.append({
                        'import_employee_id': emp_code,
                        'import_date': target_date,
                        'is_week_off': True,
                    })
                    continue

                shift_code = str(shift_code).strip()
                calendar = self.env['resource.calendar'].search([('name', '=', shift_code)], limit=1)
                if not calendar:
                    raise UserError(
                        f"HALTED: Shift Template '{shift_code}' does not exist in Odoo.")

                slots_to_create.append({
                    'import_employee_id': emp_code,
                    'calendar_id': calendar.id,
                    'import_date': target_date,
                })

        if not slots_to_create:
            raise UserError("No valid shifts found to import.")

        return slots_to_create

    def action_import_matrix(self):
        """Step 1: Parse → check holidays → show warning or import directly."""
        slots_to_create = self._parse_excel()
        holiday_map = self._get_holiday_map()

        # Find conflicts (skip week-off rows)
        conflicts = []
        seen = set()
        for slot in slots_to_create:
            if slot.get('is_week_off'):
                continue
            date = slot['import_date']
            emp = slot['import_employee_id']
            if date in holiday_map:
                key = (emp, date)
                if key not in seen:
                    seen.add(key)
                    conflicts.append({
                        'wizard_id': self.id,
                        'employee_code': emp,
                        'shift_date': date,
                        'holiday_name': holiday_map[date],
                    })

        if conflicts and not self.confirmed:
            # Save conflicts to warning lines and reopen with warning
            self.holiday_warning_ids.unlink()
            self.env['planning.matrix.import.warning'].create(conflicts)
            self.confirmed = True  # next click will bypass check

            return {
                'type': 'ir.actions.act_window',
                'res_model': 'planning.matrix.import',
                'res_id': self.id,
                'view_mode': 'form',
                'view_id': self.env.ref(
                    'attendance_planning.view_planning_matrix_import_warning_form'
                ).id,
                'target': 'new',
            }

        # No conflicts OR user confirmed → do the import
        return self._do_import(slots_to_create)

    def action_confirm_import(self):
        """Step 2: User clicked 'Yes, Import Anyway' on warning screen."""
        slots_to_create = self._parse_excel()
        return self._do_import(slots_to_create)

    def _do_import(self, slots_to_create):
        self.env['planning.slot'].create(slots_to_create)
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Success',
                'message': f'Successfully imported {len(slots_to_create)} shifts!',
                'type': 'success',
                'sticky': False,
            }
        }


class PlanningMatrixImportWarning(models.TransientModel):
    _name = 'planning.matrix.import.warning'
    _description = 'Holiday Conflict Warning Line'

    wizard_id = fields.Many2one('planning.matrix.import', string='Wizard')
    employee_code = fields.Char(string='Employee ID')
    shift_date = fields.Date(string='Date')
    holiday_name = fields.Char(string='Public Holiday')